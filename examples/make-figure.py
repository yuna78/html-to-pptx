#!/usr/bin/env python3
"""Regenerate the README figure: HTML source vs the converted PPTX with shape outlines.

Usage:  .venv/bin/python examples/make-figure.py [--slide 2] [-o examples/figure-editable.png]

Needs Chrome (for the HTML screenshot), LibreOffice + poppler (to render the PPTX back
to an image) and Pillow. It is a maintenance tool, not part of the conversion pipeline.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "engine"))

PANEL_W = 1200
BG = (244, 246, 250)
INK = (16, 35, 63)
ACCENT = (47, 127, 212)
FONT_CANDIDATES = [
    # (path, face index, is_bold) — sans-serif CJK faces first; the label strip must not
    # fall back to a serif face or the figure looks like a different tool made it.
    ("/System/Library/Fonts/Hiragino Sans GB.ttc", 2, True),
    ("/System/Library/Fonts/Hiragino Sans GB.ttc", 0, False),
    ("/System/Library/Fonts/PingFang.ttc", 4, True),
    ("/System/Library/Fonts/PingFang.ttc", 2, False),
    ("/System/Library/Fonts/STHeiti Medium.ttc", 1, True),
    ("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc", 0, True),
    ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", 0, False),
    ("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc", 0, False),
]


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    for path, index, is_bold in FONT_CANDIDATES:
        if bold and not is_bold:
            continue
        if not Path(path).exists():
            continue
        try:
            return ImageFont.truetype(path, size, index=index)
        except Exception:
            continue
    for path, index, _ in FONT_CANDIDATES:            # any face beats the bitmap default
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size, index=index)
            except Exception:
                continue
    return ImageFont.load_default()


def screenshot_html(html: Path, slide_index: int, out_png: Path, size: tuple[int, int]) -> None:
    """Render one .slide of the deck in headless Chrome."""
    from html2pptx import discover_chrome

    chrome = discover_chrome()
    if not chrome:
        raise SystemExit("Chrome not found — cannot screenshot the HTML source.")
    source = html.read_text(encoding="utf-8")
    isolate = (
        "<style>body{margin:0!important;background:#fff!important}"
        f".slide:not(:nth-of-type({slide_index})){{display:none!important}}"
        ".slide{margin:0!important}</style>"
    )
    patched = source.replace("</head>", isolate + "</head>", 1)
    temp = html.with_name(f".figure-slide-{slide_index}.html")
    temp.write_text(patched, encoding="utf-8")
    try:
        subprocess.run(
            [str(chrome), "--headless=new", "--disable-gpu", "--hide-scrollbars",
             f"--window-size={size[0]},{size[1]}", f"--screenshot={out_png}", str(temp)],
            check=True, capture_output=True,
        )
    finally:
        temp.unlink(missing_ok=True)


def render_pptx(pptx: Path, slide_index: int, out_png: Path, workdir: Path, width: int) -> None:
    """Render the PPTX back to an image through LibreOffice + poppler."""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not Path(soffice).exists():
        raise SystemExit("LibreOffice not found — cannot render the PPTX.")
    subprocess.run(
        [soffice, "--headless", f"-env:UserInstallation=file://{workdir}/profile",
         "--convert-to", "pdf", "--outdir", str(workdir), str(pptx)],
        check=True, capture_output=True,
    )
    pdf = workdir / f"{pptx.stem}.pdf"
    subprocess.run(
        ["pdftoppm", "-png", "-scale-to-x", str(width), "-scale-to-y", "-1",
         "-f", str(slide_index), "-l", str(slide_index), str(pdf), str(workdir / "page")],
        check=True, capture_output=True,
    )
    rendered = sorted(workdir.glob("page-*.png"))[0]
    shutil.copy(rendered, out_png)


def shape_boxes(pptx: Path, slide_index: int) -> list[tuple[float, float, float, float]]:
    """Every leaf shape's box, in slide-relative fractions."""
    from pptx import Presentation

    presentation = Presentation(str(pptx))
    slide = presentation.slides[slide_index - 1]
    width, height = presentation.slide_width, presentation.slide_height
    boxes: list[tuple[float, float, float, float]] = []

    def walk(shapes):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                walk(shape.shapes)
                continue
            if None in (shape.left, shape.top, shape.width, shape.height):
                continue
            boxes.append((shape.left / width, shape.top / height,
                          shape.width / width, shape.height / height))

    walk(slide.shapes)
    return boxes


def panel(image: Image.Image, caption: str, sub: str | None = None) -> Image.Image:
    """One labelled panel: caption strip above a bordered screenshot."""
    scaled = image.resize((PANEL_W, round(image.height * PANEL_W / image.width)), Image.LANCZOS)
    strip = 46
    canvas = Image.new("RGB", (PANEL_W + 4, scaled.height + strip + 4), BG)
    draw = ImageDraw.Draw(canvas)
    draw.ellipse((2, strip // 2 - 7, 16, strip // 2 + 7), fill=ACCENT)
    draw.text((26, strip // 2 - 12), caption, font=load_font(20, bold=True), fill=INK)
    if sub:
        title_width = draw.textlength(caption, font=load_font(20, bold=True))
        draw.text((26 + title_width + 14, strip // 2 - 9), sub, font=load_font(16), fill=(112, 134, 163))
    canvas.paste(scaled, (2, strip))
    draw.rectangle((1, strip - 1, PANEL_W + 2, strip + scaled.height), outline=(206, 216, 232))
    return canvas


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--deck", type=Path, default=ROOT / "examples/showcase-deck.html")
    parser.add_argument("--slide", type=int, default=2)
    parser.add_argument("-o", "--output", type=Path, default=ROOT / "examples/figure-editable.png")
    args = parser.parse_args()

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        pptx = work / "deck.pptx"
        subprocess.run(
            [sys.executable, str(ROOT / "convert.py"), str(args.deck), "-o", str(pptx), "--quiet"],
            check=True,
        )

        html_png = work / "html.png"
        pptx_png = work / "pptx.png"
        screenshot_html(args.deck, args.slide, html_png, (1280, 720))
        render_pptx(pptx, args.slide, pptx_png, work, PANEL_W)

        outlined = Image.open(pptx_png).convert("RGB")
        draw = ImageDraw.Draw(outlined, "RGBA")
        boxes = shape_boxes(pptx, args.slide)
        for x, y, w, h in boxes:
            draw.rectangle(
                (x * outlined.width, y * outlined.height,
                 (x + w) * outlined.width, (y + h) * outlined.height),
                outline=(47, 127, 212, 150), width=2,
            )

        top = panel(Image.open(html_png).convert("RGB"), "① 输入：浏览器里的 HTML")
        bottom = panel(outlined, "② 输出：PowerPoint 里的 PPTX",
                       f"蓝框 = {len(boxes)} 个可以直接改的原生形状")

        gap = 22
        figure = Image.new("RGB", (top.width, top.height + bottom.height + gap), BG)
        figure.paste(top, (0, 0))
        figure.paste(bottom, (0, top.height + gap))
        args.output.parent.mkdir(parents=True, exist_ok=True)
        figure.save(args.output, optimize=True)
        print(f"{args.output}  ({figure.width}x{figure.height}, {len(boxes)} shapes outlined)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
