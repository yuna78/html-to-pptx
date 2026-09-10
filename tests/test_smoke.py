"""End-to-end smoke test: needs Node.js and Chrome/Chromium on PATH."""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "engine"))


def _has_engine_prereqs() -> bool:
    import shutil

    from html2pptx import discover_chrome

    return bool(shutil.which("node")) and bool(discover_chrome())


requires_engine = pytest.mark.skipif(
    not _has_engine_prereqs(), reason="Node.js and Chrome are required for the smoke test"
)


def _shape_text(pptx_path: Path) -> list[str]:
    from pptx import Presentation

    def walk(shapes, out):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                walk(shape.shapes, out)
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())

    presentation = Presentation(str(pptx_path))
    collected: list[str] = []
    for slide in presentation.slides:
        walk(slide.shapes, collected)
    return collected


def _convert(source: Path, output: Path) -> None:
    result = subprocess.run(
        [sys.executable, str(ROOT / "convert.py"), str(source), "-o", str(output)],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert output.exists()


@requires_engine
def test_sample_deck_converts_to_editable_shapes(tmp_path):
    from pptx import Presentation

    output = tmp_path / "sample-deck.pptx"
    _convert(ROOT / "examples" / "sample-deck.html", output)

    presentation = Presentation(str(output))
    assert len(presentation.slides) == 1
    assert presentation.slide_width == 1280 * 9525
    texts = _shape_text(output)
    assert any("html2pptx" in t or "Editable" in t for t in texts), texts


@requires_engine
def test_wide_canvas_is_detected(tmp_path):
    from pptx import Presentation

    deck = tmp_path / "wide.html"
    deck.write_text(
        "<!doctype html><html><head><style>.slide{width:1600px;height:900px;background:#fff}"
        "</style></head><body><section class='slide'><h1>Wide</h1></section></body></html>",
        encoding="utf-8",
    )
    output = tmp_path / "wide.pptx"
    _convert(deck, output)

    presentation = Presentation(str(output))
    assert presentation.slide_width == 1600 * 9525
    assert presentation.slide_height == 900 * 9525
    assert "Wide" in _shape_text(output)
